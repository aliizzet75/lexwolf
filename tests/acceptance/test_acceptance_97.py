import os
import sys

from docx import Document

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '../../backend'))

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), '../../backend/templates/docx')

TYPES = [
    'klage', 'klageerwiderung', 'widerspruch',
    'antrag', 'anwaltsschreiben', 'mahnschreiben',
]

REQUIRED_PLACEHOLDERS = {
    'klage': ['{ANTRAEGE}', '{BEGRUENDUNG}', '{ANWALT}'],
    'klageerwiderung': ['{ANTRAEGE}', '{BEGRUENDUNG}', '{ANWALT}'],
    'widerspruch': ['{ANTRAEGE}', '{BEGRUENDUNG}', '{ANWALT}'],
    'antrag': ['{ANTRAEGE}', '{BEGRUENDUNG}', '{ANWALT}'],
    'anwaltsschreiben': ['{BEGRUENDUNG}', '{ANWALT}'],
    'mahnschreiben': ['{BEGRUENDUNG}', '{ANWALT}'],
}


def test_all_six_templates_exist():
    for typ in TYPES:
        path = os.path.join(TEMPLATE_DIR, f'{typ}.docx')
        assert os.path.exists(path), f'Template fehlt: {path}'


def test_templates_contain_required_structure_placeholders():
    for typ in TYPES:
        doc = Document(os.path.join(TEMPLATE_DIR, f'{typ}.docx'))
        full_text = '\n'.join(p.text for p in doc.paragraphs)
        for placeholder in REQUIRED_PLACEHOLDERS[typ]:
            assert placeholder in full_text, f'{typ}.docx fehlt Platzhalter {placeholder}'


def test_export_route_selects_template_by_schriftsatz_typ():
    from api.routes.export import _load_template

    for typ in TYPES:
        document = _load_template(typ)
        assert document is not None, f'Template konnte nicht geladen werden: {typ}'
